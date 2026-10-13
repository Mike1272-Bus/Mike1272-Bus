import sys, urllib.parse
sys.path.insert(0, '/tmp/claude-0/-home-user-Mike1272-Bus/7c68b9b5-f98d-5769-9dfc-b4f40cf3c57b/scratchpad')
from prospects_data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = sys.argv[1]
F = 'Arial'; NAVY = '1F3864'
IN = PatternFill('solid', fgColor='FFF2CC'); AUTO = PatternFill('solid', fgColor='EDEDED'); HEAD = PatternFill('solid', fgColor=NAVY)
th = Side(style='thin', color='BFBFBF'); BD = Border(left=th, right=th, top=th, bottom=th)
def fn(**k): k.setdefault('name', F); k.setdefault('size', 10); return Font(**k)
STATUTS = ['À contacter', 'Message envoyé', 'Relance 1 faite', 'Relance 2 faite', 'A répondu', 'Démo prévue', 'Démo faite', 'Client', 'Pas intéressé', 'Sans réponse']

rows = sorted(load(), key=lambda r: (vague(r), SEG_ORDER.index(r['seg']), r['ville'], r['n']))
wb = Workbook(); ws = wb.active; ws.title = 'Prospects'
ws['A1'] = 'Prospection : template « Gestion de stock » et système automatisé'; ws['A1'].font = fn(size=15, bold=True, color=NAVY)
ws['A2'] = 'Jaune = à remplir · Gris = automatique · Cliquez sur « Ouvrir WhatsApp » : le message 1 est déjà écrit, ajoutez la vidéo puis envoyez.'
ws['A2'].font = fn(size=9, italic=True, color='595959')
cols = [('N°', 5), ('Vague', 7), ('Priorité', 8), ('Entreprise', 30), ('Ville', 12), ('Téléphone', 18), ('WhatsApp', 16),
        ('Type (fichier)', 22), ('Segment', 26), ('Date du 1er message', 13), ('Statut', 16), ('Relance 1 le', 12), ('Relance 2 le', 12),
        ('Prochaine action', 30), ('Réponse / notes', 36), ('À vérifier', 44)]
HR = 4
for i, (h, w) in enumerate(cols, 1):
    c = ws.cell(HR, i, h); c.font = fn(bold=True, color='FFFFFF'); c.fill = HEAD; c.border = BD
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[HR].height = 30
for k, r in enumerate(rows):
    y = HR + 1 + k
    url = 'https://wa.me/' + r['wa'] + '?text=' + urllib.parse.quote(message1(r['nom'], r['seg']))
    vals = [r['n'], vague(r), r['prio'], r['nom'], r['ville'], r['tel'], 'Ouvrir WhatsApp', r['type_orig'], r['seg'], None, 'À contacter',
            f'=IF(J{y}="","",J{y}+3)', f'=IF(J{y}="","",J{y}+7)',
            (f'=IF(K{y}="À contacter","Envoyer message 1 + vidéo présentation",'
             f'IF(K{y}="Message envoyé",IF(J{y}="","Indiquer la date d\'envoi",IF(TODAY()>=J{y}+3,"Envoyer la relance 1","Attendre la réponse")),'
             f'IF(K{y}="Relance 1 faite",IF(TODAY()>=J{y}+7,"Envoyer la relance 2","Attendre la réponse"),'
             f'IF(K{y}="Relance 2 faite","Si rien sous 3 jours : Sans réponse",'
             f'IF(K{y}="A répondu","Envoyer message 2 + vidéo workflow",'
             f'IF(K{y}="Démo prévue","Préparer la démo avec ses produits",'
             f'IF(K{y}="Démo faite","Envoyer l\'offre et le prix",'
             f'IF(K{y}="Client","Installer et demander un avis","Fini"))))))))'),
            None, r['notes']]
    for i, v in enumerate(vals, 1):
        c = ws.cell(y, i, v); c.font = fn(); c.border = BD
        c.fill = IN if i in (10, 11, 15) else (AUTO if i in (12, 13, 14) else PatternFill())
        c.alignment = Alignment(vertical='center', wrap_text=i in (14, 15, 16))
    lk = ws.cell(y, 7); lk.hyperlink = url; lk.font = fn(color='0563C1', underline='single')
    for i in (10, 12, 13): ws.cell(y, i).number_format = 'DD/MM/YYYY'
    for i in (1, 2, 3): ws.cell(y, i).alignment = Alignment(horizontal='center', vertical='center')
last = HR + len(rows)
dv = DataValidation(type='list', formula1='"' + ','.join(STATUTS) + '"', allow_blank=False, showErrorMessage=True, error='Choisissez un statut dans la liste.')
ws.add_data_validation(dv); dv.add(f'K{HR+1}:K{last}')
dvd = DataValidation(type='date', operator='greaterThan', formula1='45000', allow_blank=True); ws.add_data_validation(dvd); dvd.add(f'J{HR+1}:J{last}')
for st, col in [('Client', 'C6EFCE'), ('Démo prévue', 'DDEBF7'), ('Démo faite', 'DDEBF7'), ('A répondu', 'FFEB9C'), ('Pas intéressé', 'F4B6B6'), ('Sans réponse', 'E7E6E6')]:
    ws.conditional_formatting.add(f'K{HR+1}:K{last}', CellIsRule(operator='equal', formula=[f'"{st}"'], fill=PatternFill('solid', fgColor=col)))
for p_, col in [('A', 'C6EFCE'), ('B', 'FFEB9C'), ('C', 'F4B6B6')]:
    ws.conditional_formatting.add(f'C{HR+1}:C{last}', CellIsRule(operator='equal', formula=[f'"{p_}"'], fill=PatternFill('solid', fgColor=col)))
ws.conditional_formatting.add(f'N{HR+1}:N{last}', FormulaRule(formula=[f'LEFT(N{HR+1},7)="Envoyer"'], font=Font(name=F, bold=True, color='9C0006')))
ws.freeze_panes = f'E{HR+1}'; ws.auto_filter.ref = f'A{HR}:P{last}'

# --- Suivi
su = wb.create_sheet('Suivi', 0)
su.column_dimensions['A'].width = 3; su.column_dimensions['B'].width = 34; su.column_dimensions['C'].width = 12
for c in 'DEFG': su.column_dimensions[c].width = 12
su['B1'] = 'Suivi de la prospection'; su['B1'].font = fn(size=16, bold=True, color=NAVY)
K = f'Prospects!$K${HR+1}:$K${last}'; V = f'Prospects!$B${HR+1}:$B${last}'
su['B3'] = 'Statut'; su['C3'] = 'Total'; su['D3'] = 'Vague 1'; su['E3'] = 'Vague 2'; su['F3'] = 'Vague 3'
for c in ('B3', 'C3', 'D3', 'E3', 'F3'): su[c].font = fn(bold=True, color='FFFFFF'); su[c].fill = HEAD; su[c].border = BD
for i, st in enumerate(STATUTS):
    y = 4 + i
    su.cell(y, 2, st).border = BD
    su.cell(y, 3, f'=COUNTIF({K},B{y})').border = BD
    for j, v in enumerate((1, 2, 3)):
        su.cell(y, 4 + j, f'=COUNTIFS({K},$B{y},{V},{v})').border = BD
    for x in range(2, 7): su.cell(y, x).font = fn()
y = 4 + len(STATUTS)
su.cell(y, 2, 'Total').font = fn(bold=True); su.cell(y, 3, f'=SUM(C4:C{y-1})').font = fn(bold=True)
y += 2
kp = [('Contactés (au moins 1 message)', f'=C4+0', None)]
su.cell(y, 2, 'Indicateurs').font = fn(size=12, bold=True, color=NAVY); y += 1
ind = [('Entreprises contactées', f'=C{4+len(STATUTS)}-C4', '0'),
       ('Ont répondu (ou plus loin)', '=C8+C9+C10+C11', '0'),
       ('Taux de réponse', f'=IF(C{y}=0,"-",C{y+1}/C{y})', '0%'),
       ('Démos faites', '=C10+C11', '0'),
       ('Clients', '=C11', '0')]
for lab, fml, nf in ind:
    su.cell(y, 2, lab).font = fn(); su.cell(y, 3, fml).font = fn(bold=True); su.cell(y, 3).number_format = nf
    su.cell(y, 2).border = BD; su.cell(y, 3).border = BD; su.cell(y, 3).fill = AUTO; y += 1
y += 1
su.cell(y, 2, 'Le taux de réponse compte « A répondu », « Démo prévue », « Démo faite » et « Client ».').font = fn(size=9, italic=True, color='7F7F7F')

# --- Messages
ms = wb.create_sheet('Messages')
ms.column_dimensions['A'].width = 26; ms.column_dimensions['B'].width = 110
ms['A1'] = 'Les messages (vous)'; ms['A1'].font = fn(size=15, bold=True, color=NAVY)
items = [('Message 1 (jour 0) + vidéo de présentation', message1('[Nom]', 'Téléphones, électronique, informatique').replace('\\n', '\n')),
         ('Phrase d\'accroche selon le segment', '\n'.join(f'• {s} : {h}' for s, h in HOOK.items())),
         ('Message 2 (s\'il répond) + vidéo du workflow', MESSAGE2.replace('\\n', '\n')),
         ('Relance 1 (jour 3)', RELANCE1.replace('{nom}', '[Nom]')),
         ('Relance 2 (jour 7, la dernière)', RELANCE2.replace('{nom}', '[Nom]'))]
for i, (a, b) in enumerate(items):
    y = 3 + i
    ms.cell(y, 1, a).font = fn(bold=True); ms.cell(y, 2, b).font = fn()
    for c in (1, 2): ms.cell(y, c).alignment = Alignment(wrap_text=True, vertical='top'); ms.cell(y, c).border = BD
    ms.row_dimensions[y].height = max(45, 15 * (b.count('\n') + 2 + len(b) // 110))
wb.save(OUT); print('ok', len(rows))
