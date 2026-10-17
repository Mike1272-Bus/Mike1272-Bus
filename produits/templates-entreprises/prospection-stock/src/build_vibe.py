import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
SRC = sys.argv[1]; OUT = sys.argv[2]
rows = list(csv.DictReader(open(SRC, encoding='utf-8')))
J, C, D, X = 'Joignable : ajoutée au fichier de suivi', 'Pertinente : contact à chercher', 'Déjà dans ta prospection', 'Écartée'
V = {  # n -> (verdict, activité, raison, contact, source)
 1: (X, 'Brasserie (Bralima)', 'Très grande entreprise : déjà équipée de logiciels', '', ''),
 2: (X, 'Boissons (Varun, PepsiCo)', 'Très grande entreprise', '', ''),
 3: (X, 'Plateforme agricole', 'Pas de stock de produits à vendre', '', ''),
 4: (X, 'Sucrerie (CSKN)', 'Très grande entreprise', '', ''),
 5: (D, 'Café La Kinoise', 'Déjà contactée pour la vidéo animée', '', ''),
 6: (X, 'Minoterie (farines)', 'Doublon de la ligne 40 (MIDEMA)', '', ''),
 7: (X, 'Boissons (Refriango)', 'Très grande entreprise', '', ''),
 8: (C, 'Chocolat, café, savon (Virunga)', 'Produits physiques, boutique en ligne', 'aucun sur le site', 'origins.virunga.org'),
 9: (D, 'Transformation alimentaire', 'Déjà contactée (ManiTech)', '', ''),
 10: (J, 'Légumes déshydratés', 'Vend aux ménages et aux commerces', '+260 970 463 609 · umandesarl@gmail.com', 'Leur site Canva'),
 11: (X, 'Eau (Swissta)', 'Grande entreprise', '', ''),
 12: (D, 'Vins Tomela', 'Déjà contactée', '', ''),
 13: (X, 'Agro-industrie (Mashamba)', 'Grande entreprise', '', ''),
 14: (X, 'Boissons (Kabeco, Bora)', 'Grande entreprise', '', ''),
 15: (D, 'Miel (La Ruche du Terroir)', 'Déjà en semaine test', '', ''),
 16: (X, 'Eau (Swissta)', 'Doublon, grande entreprise', '', ''),
 17: (D, 'Pâtisserie', 'Déjà contactée (Délices d\'Émilie)', '', ''),
 18: (C, 'Huile de palme', 'Vend en bidons de 1 et 5 L (via JauneCongo)', 'aucun trouvé', 'Facebook / Instagram de JauneCongo'),
 19: (D, 'Boissons Festa', 'Déjà contactée', '', ''),
 20: (C, 'Jus naturels (Goma)', 'Production et distribution de boissons', 'aucun sur le site', 'kivusocietycorporation.com'),
 21: (J, 'Farines de bouillie (Bukavu)', 'Petite entreprise, produits en sachets', '+243 997 374 225 (à vérifier)', 'Site ButamuCare'),
 22: (C, 'Apiculture, miel', 'Produits physiques', 'site hors ligne', ''),
 23: (X, 'Distribution (Liberty)', 'Très grande entreprise', '', ''),
 24: (C, 'Le Tablier (restaurant ?)', 'Activité peu claire', 'page Facebook', 'facebook.com/letablierrdc'),
 25: (C, 'Tisanes Yanzambi', 'Produits emballés, points de vente', 'à chercher sur le site', 'yanzambi.com'),
 26: (X, 'Hôtellerie, lounges', 'Groupe de restaurants, besoin différent', '', ''),
 27: (D, 'Épicé Bon (assaisonnements)', 'Déjà contactée. Numéro : +243 816 757 584', '', ''),
 28: (J, 'Réseau de boucheries (Lubumbashi)', 'Stock et chaîne du froid, plusieurs points de vente', '+243 997 022 378 · info@frigocongo.com', 'Page Facebook FrigoCongo'),
 29: (X, 'Industrie (savon, huile, boissons)', 'Très grande entreprise', '', ''),
 30: (X, 'Import et distribution (Sokin)', 'Très grande entreprise (23 points de vente)', '', ''),
 31: (X, 'Export café et cacao', 'Pas de vente au détail', '', ''),
 32: (C, 'Frites surgelées (FRICO)', 'Production et vente', 'site hors ligne', 'Fondateur : Jean Johnson Bapanga (VC4A)'),
 33: (C, 'Chawarma (restaurant mobile)', 'Stock d\'ingrédients', 'page Facebook', 'facebook.com/chawarpips'),
 34: (C, 'Boulangerie', 'Stock d\'ingrédients, pas de contact', 'aucun', ''),
 35: (X, 'Agriculture (TFA)', 'Activité trop vague', '', ''),
 36: (D, 'Kahawa Congo Coffee', 'Déjà contactée', '', ''),
 37: (X, 'Services (restaurant, nettoyage, fleurs)', 'Activité trop dispersée', '', ''),
 38: (C, 'Eau en bouteille (Lubumbashi)', 'Production et distribution', 'pas de numéro public', '5393 Av. des Cimetières, Lubumbashi (pearl-industries.com)'),
 39: (X, 'Agroforesterie', 'Pas de vente de produits en boutique', '', ''),
 40: (J, 'Minoterie (farines, aliments bétail)', 'Stock important. Peut-être déjà équipée', '+243 81 555 8000', 'Site Seaboard (standard)'),
 41: (X, 'Application de livraison', 'Pas de stock', '', ''),
 42: (C, 'Livraison de repas (KenBen)', 'Site hors ligne', 'aucun', ''),
 43: (X, 'Restaurant gastronomique', 'Besoin différent', '', ''),
 44: (J, 'Livraison produits alimentaires', 'Stock de poissons et viandes, plusieurs villes', '+243 895 895 888 (WhatsApp) · +243 81 222 9100 · biofoodscongo@gmail.com', 'Leur description'),
 45: (C, 'Agroalimentaire (Lubumbashi)', 'Projet encore jeune', 'aucun', 'LinkedIn'),
 46: (X, 'Initiative communautaire', 'Pas de boutique', '', ''),
 47: (C, 'Gingembre en poudre (Bukavu)', 'Produit emballé', 'aucun', 'LinkedIn'),
 48: (C, 'Agro bio (Goma)', 'Activité vague', 'aucun', ''),
 49: (X, 'Restaurants (franchise PAUL)', 'Groupe déjà structuré', '', ''),
 50: (X, 'Activités multiples (MONGO Group)', 'Activité trop dispersée', '', ''),
}
wb = Workbook(); ws = wb.active; ws.title = 'Analyse'
th = Side(style='thin', color='BFBFBF'); BD = Border(left=th, right=th, top=th, bottom=th)
ws['A1'] = 'Liste Vibe Prospection (50 entreprises alimentaires) : analyse pour le système de stock'; ws['A1'].font = Font(name='Arial', size=14, bold=True, color='1F3864')
ws['A2'] = 'Contacts cherchés le 13/10/2026 sur leurs sites et pages publiques. À vérifier avant d\'écrire.'; ws['A2'].font = Font(name='Arial', size=9, italic=True)
cols = [('N°', 5), ('Entreprise', 32), ('Ville', 13), ('Employés', 11), ('Activité', 30), ('Verdict', 30), ('Pourquoi', 40), ('Contact trouvé', 38), ('Source', 32)]
for i, (h, w) in enumerate(cols, 1):
    c = ws.cell(4, i, h); c.font = Font(name='Arial', bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.border = BD
    ws.column_dimensions[c.column_letter].width = w
order = {J: 0, C: 1, D: 2, X: 3}
data = sorted(rows, key=lambda r: (order[V[int(r['row_num'])][0]], int(r['row_num'])))
for k, r in enumerate(data):
    n = int(r['row_num']); v = V[n]; y = 5 + k
    vals = [n, r['business_name'], (r['business_city_name'] or '?').title(), r['business_number_of_employees_range'].strip('[]'), v[1], v[0], v[2], v[3], v[4]]
    for i, val in enumerate(vals, 1):
        c = ws.cell(y, i, val); c.font = Font(name='Arial', size=10); c.border = BD; c.alignment = Alignment(wrap_text=True, vertical='top')
last = 4 + len(data)
for txt, col in [(J, 'C6EFCE'), (C, 'FFEB9C'), (D, 'DDEBF7'), (X, 'EDEDED')]:
    ws.conditional_formatting.add(f'F5:F{last}', CellIsRule(operator='equal', formula=[f'"{txt}"'], fill=PatternFill('solid', fgColor=col)))
ws.freeze_panes = 'C5'; ws.auto_filter.ref = f'A4:I{last}'
from collections import Counter
cnt = Counter(V[int(r['row_num'])][0] for r in rows)
ws['K4'] = 'Résumé'; ws['K4'].font = Font(name='Arial', bold=True)
for i, t in enumerate([J, C, D, X]):
    ws.cell(5 + i, 11, t).font = Font(name='Arial'); ws.cell(5 + i, 12, f'=COUNTIF(F5:F{last},K{5+i})').font = Font(name='Arial', bold=True)
ws.column_dimensions['K'].width = 36
wb.save(OUT); print(dict(cnt))
