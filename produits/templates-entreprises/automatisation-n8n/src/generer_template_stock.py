import datetime, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

OUT = sys.argv[1]
AUTO = "--auto" in sys.argv
C_LAST = 5000
CM = "Commandes!$%s$2:$%s$" + str(C_LAST)
AR = "Arrivages!$%s$2:$%s$" + str(C_LAST)
def arr_sum(crit):  # quantités reçues par le formulaire d'arrivage (n8n)
    if not AUTO: return ""
    return f'+SUMIFS({AR % ("D","D")},{AR % ("B","B")},{crit})'
def cmd_sum(col, crit_col, crit):  # somme des commandes n8n (hors annulées)
    if not AUTO: return ""
    return f'+SUMIFS({CM % (col, col)},{CM % (crit_col, crit_col)},{crit},{CM % ("J","J")},"<>Annulée")'
F = "Arial"
NAVY = "1F3864"
IN_FILL = PatternFill("solid", fgColor="FFF2CC")   # à remplir
AUTO_FILL = PatternFill("solid", fgColor="EDEDED")  # automatique
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;-#,##0;"-"'
P_FIRST, P_LAST = 5, 204
M_FIRST, M_LAST = 5, 1004

def font(**k):
    k.setdefault("name", F); k.setdefault("size", 10)
    return Font(**k)

def title(ws, text, sub, width):
    ws["A1"] = text; ws["A1"].font = font(size=16, bold=True, color=NAVY)
    ws["A2"] = sub; ws["A2"].font = font(size=9, italic=True, color="595959")
    ws.row_dimensions[1].height = 26

def headers(ws, row, cols):
    for i, (name, w, kind) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32

wb = Workbook()

# ---------- Mode d'emploi ----------
me = wb.active; me.title = "Mode d'emploi"
me.column_dimensions["A"].width = 4; me.column_dimensions["B"].width = 95
me["B1"] = "Gestion de stock : mode d'emploi"; me["B1"].font = font(size=18, bold=True, color=NAVY)
lines = [
    ("h", "À quoi sert ce tableau"),
    ("t", "Vous savez à tout moment ce qu'il vous reste en stock, quels produits commander, ce que vaut votre stock, et ce que vous avez vendu et gagné ce mois-ci."),
    ("h", "Les couleurs"),
    ("in", "Cases jaunes : vous les remplissez."),
    ("auto", "Cases grises : elles se calculent toutes seules. N'écrivez rien dedans."),
    ("h", "Étape 1 : une seule fois, au départ (onglet Produits)"),
    ("t", "1. Pour chaque produit : un code court (ex. P001), le nom, la catégorie, l'unité (pièce, carton, kg…)."),
    ("t", "2. Le prix d'achat et le prix de vente d'une unité."),
    ("t", "3. Le stock de départ : ce que vous comptez aujourd'hui sur vos étagères."),
    ("t", "4. Le seuil d'alerte : en dessous de cette quantité, le produit passe en « À commander »."),
    ("t", "5. Facultatif : le fournisseur et son téléphone. La colonne « Quantité à commander » propose de quoi remonter le stock au double du seuil d'alerte."),
    ("h", "Étape 2 : chaque jour (onglet Mouvements)"),
    ("t", "Une ligne par mouvement : la date, le code du produit (liste déroulante), Entrée (vous recevez de la marchandise) ou Sortie (vous vendez), et la quantité."),
    ("t", "Le nom du produit, le prix et le montant se remplissent tout seuls."),
    ("h", "Étape 3 : regarder le résultat (onglet Tableau de bord)"),
    ("t", "Choisissez la devise et le mois à afficher. Vous voyez : la valeur du stock, les produits à commander ou en rupture, les ventes, les achats et la marge du mois."),
    ("t", "Dans l'onglet Produits, la colonne État devient orange (À commander) ou rouge (Rupture)."),
    ("h", "Avant de commencer"),
    ("t", "Les lignes marquées EXEMPLE montrent comment remplir. Supprimez-les (produits et mouvements) avant de saisir vos vrais chiffres."),
    ("t", "Capacité : 200 produits et 1 000 mouvements. Au-delà, demandez une version plus grande."),
    ("t", "Astuce : le tableau marche aussi sur téléphone, dans Google Sheets ou l'application Excel."),
]
r = 3
for kind, txt in lines:
    c = me.cell(row=r, column=2, value=txt)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "h":
        c.font = font(size=12, bold=True, color=NAVY); r += 1
        me.row_dimensions[r-1].height = 22
    else:
        c.font = font(size=10)
        if kind == "in": c.fill = IN_FILL
        if kind == "auto": c.fill = AUTO_FILL
    r += 1
me.sheet_view.showGridLines = False

# ---------- Produits ----------
pr = wb.create_sheet("Produits")
title(pr, "Produits", "Jaune = à remplir · Gris = calculé automatiquement · Supprimez les lignes EXEMPLE avant de commencer", 13)
pcols = [("Code", 10, "in"), ("Produit", 28, "in"), ("Catégorie", 16, "in"), ("Unité", 10, "in"),
         ("Prix d'achat (1 unité)", 13, "in"), ("Prix de vente (1 unité)", 13, "in"),
         ("Stock de départ", 11, "in"), ("Seuil d'alerte", 10, "in"),
         ("Entrées", 10, "auto"), ("Sorties", 10, "auto"), ("Stock actuel", 11, "auto"),
         ("Valeur du stock (prix d'achat)", 15, "auto"), ("État", 14, "auto"),
         ("Fournisseur", 20, "in"), ("Tél. fournisseur", 16, "in"), ("Quantité à commander", 13, "auto")]
headers(pr, 4, pcols)
MV = f"Mouvements!$%s${M_FIRST}:$%s${M_LAST}"
for r in range(P_FIRST, P_LAST + 1):
    f = {
        9:  f'=IF($A{r}="","",SUMIFS({MV % ("E","E")},{MV % ("B","B")},$A{r},{MV % ("D","D")},"Entrée"){arr_sum("$A"+str(r))})',
        16: f'=IF($A{r}="","",IF($K{r}<=N($H{r}),MAX(0,2*N($H{r})-$K{r}),0))',
        10: f'=IF($A{r}="","",SUMIFS({MV % ("E","E")},{MV % ("B","B")},$A{r},{MV % ("D","D")},"Sortie"){cmd_sum("F", "D", "$A"+str(r))})',
        11: f'=IF($A{r}="","",N($G{r})+$I{r}-$J{r})',
        12: f'=IF($A{r}="","",$K{r}*N($E{r}))',
        13: f'=IF($A{r}="","",IF($K{r}<=0,"Rupture",IF($K{r}<=N($H{r}),"À commander","OK")))',
    }
    for ci, (_, _, kind) in enumerate(pcols, start=1):
        c = pr.cell(row=r, column=ci)
        if ci in f: c.value = f[ci]
        c.fill = IN_FILL if kind == "in" else AUTO_FILL
        c.border = BORDER; c.font = font()
        if ci in (5, 6, 12): c.number_format = NUM
        if ci in (7, 8, 9, 10, 11, 16): c.number_format = NUM
        if ci == 13: c.alignment = Alignment(horizontal="center")
ex_p = [("P001", "EXEMPLE - Riz 25 kg", "Alimentation", "sac", 30, 36, 20, 5) + (None,)*5 + ("EXEMPLE - Grossiste A",),
        ("P002", "EXEMPLE - Huile 5 L", "Alimentation", "bidon", 12, 15, 10, 4) + (None,)*5 + ("EXEMPLE - Grossiste B",)]
for i, row in enumerate(ex_p):
    for ci, v in enumerate(row, start=1):
        if v is None: continue
        pr.cell(row=P_FIRST + i, column=ci, value=v).font = font(italic=True, color="7F7F7F")
rng = f"M{P_FIRST}:M{P_LAST}"
pr.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Rupture"'], fill=PatternFill("solid", fgColor="F4B6B6"), font=Font(name=F, bold=True, color="9C0006")))
pr.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"À commander"'], fill=PatternFill("solid", fgColor="FCD9A8"), font=Font(name=F, bold=True, color="9C5700")))
pr.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, color="006100")))
dv_dup = DataValidation(type="custom", formula1=f"COUNTIF($A${P_FIRST}:$A${P_LAST},A{P_FIRST})=1",
                        showErrorMessage=True, errorTitle="Code en double", error="Ce code existe déjà. Chaque produit doit avoir un code unique.")
pr.add_data_validation(dv_dup); dv_dup.add(f"A{P_FIRST}:A{P_LAST}")
pr.freeze_panes = "C5"

# ---------- Mouvements ----------
mv = wb.create_sheet("Mouvements")
title(mv, "Mouvements de stock", "Une ligne par entrée (réception de marchandise) ou sortie (vente). Jaune = à remplir · Gris = automatique", 9)
mcols = [("Date", 12, "in"), ("Code produit", 12, "in"), ("Produit", 28, "auto"), ("Type", 11, "in"),
         ("Quantité", 10, "in"), ("Prix unitaire", 12, "auto"), ("Montant", 13, "auto"),
         ("Coût d'achat (pour la marge)", 14, "auto"), ("Note (client, fournisseur…)", 28, "in")]
headers(mv, 4, mcols)
PA = f"Produits!$A${P_FIRST}:$A${P_LAST}"
def look(col, r): return f"INDEX(Produits!${col}${P_FIRST}:${col}${P_LAST},MATCH($B{r},{PA},0))"
for r in range(M_FIRST, M_LAST + 1):
    f = {
        3: f'=IF($B{r}="","",IFERROR({look("B", r)},"Code inconnu"))',
        6: f'=IF(OR($B{r}="",$D{r}=""),"",IFERROR(IF($D{r}="Entrée",{look("E", r)},{look("F", r)}),""))',
        7: f'=IF(OR($F{r}="",$E{r}=""),"",$E{r}*$F{r})',
        8: f'=IF(AND($D{r}="Sortie",$B{r}<>"",$E{r}<>""),IFERROR($E{r}*{look("E", r)},""),"")',
    }
    for ci, (_, _, kind) in enumerate(mcols, start=1):
        c = mv.cell(row=r, column=ci)
        if ci in f: c.value = f[ci]
        c.fill = IN_FILL if kind == "in" else AUTO_FILL
        c.border = BORDER; c.font = font()
        if ci == 1: c.number_format = "DD/MM/YYYY"
        if ci in (5, 6, 7, 8): c.number_format = NUM
today = datetime.date.today()
ex_m = [(today, "P001", "Entrée", 10, "EXEMPLE - livraison fournisseur"),
        (today, "P001", "Sortie", 4, "EXEMPLE - vente"),
        (today, "P002", "Sortie", 7, "EXEMPLE - vente")]
for i, (d, code, t, q, note) in enumerate(ex_m):
    r = M_FIRST + i
    for col, v in ((1, d), (2, code), (4, t), (5, q), (9, note)):
        mv.cell(row=r, column=col, value=v).font = font(italic=True, color="7F7F7F")
dv_type = DataValidation(type="list", formula1='"Entrée,Sortie"', allow_blank=True, showErrorMessage=True,
                         errorTitle="Type", error="Choisissez Entrée ou Sortie dans la liste.")
dv_code = DataValidation(type="list", formula1=f"={PA}", allow_blank=True, showErrorMessage=True,
                         errorTitle="Code produit", error="Ce code n'existe pas dans l'onglet Produits.")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, errorTitle="Date", error="Entrez une date (ex. 12/10/2026).")
dv_qty = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                        showErrorMessage=True, errorTitle="Quantité", error="La quantité doit être un nombre positif.")
for dv, rg in ((dv_type, "D"), (dv_code, "B"), (dv_date, "A"), (dv_qty, "E")):
    mv.add_data_validation(dv); dv.add(f"{rg}{M_FIRST}:{rg}{M_LAST}")
mv.conditional_formatting.add(f"C{M_FIRST}:C{M_LAST}", CellIsRule(operator="equal", formula=['"Code inconnu"'], font=Font(name=F, bold=True, color="C00000")))
mv.freeze_panes = "A5"

# ---------- Tableau de bord ----------
tb = wb.create_sheet("Tableau de bord", 1)
tb.sheet_view.showGridLines = False
tb.column_dimensions["A"].width = 3; tb.column_dimensions["B"].width = 38; tb.column_dimensions["C"].width = 20
tb.column_dimensions["D"].width = 50
tb["B1"] = "Tableau de bord du stock"; tb["B1"].font = font(size=18, bold=True, color=NAVY)
tb["B3"] = "Devise"; tb["C3"] = "$"
tb["B4"] = "Mois affiché (n'importe quelle date du mois)"; tb["C4"] = "=TODAY()"
tb["D3"] = "Écrivez $ ou FC"; tb["D4"] = "Remplacez par une date pour voir un autre mois"
for a in ("B3", "B4"): tb[a].font = font(bold=True)
for a in ("C3", "C4"): tb[a].fill = IN_FILL; tb[a].border = BORDER; tb[a].font = font(bold=True); tb[a].alignment = Alignment(horizontal="center")
for a in ("D3", "D4"): tb[a].font = font(size=9, italic=True, color="7F7F7F")
tb["C4"].number_format = "MMMM YYYY"
start = 'DATE(YEAR($C$4),MONTH($C$4),1)'; end = 'DATE(YEAR($C$4),MONTH($C$4)+1,1)'
def month_sum(col, typ):
    s = f'SUMIFS({MV % (col, col)},{MV % ("A","A")},">="&{start},{MV % ("A","A")},"<"&{end}'
    return s + (f',{MV % ("D","D")},"{typ}")' if typ else ")")
def cmd_month(col):
    if not AUTO: return ""
    a = CM % ("A", "A")
    return f'+SUMIFS({CM % (col, col)},{a},">="&{start},{a},"<"&{end},{CM % ("J","J")},"<>Annulée")'
def arr_month():
    if not AUTO: return ""
    a = AR % ("A", "A")
    return f'+SUMIFS({AR % ("F","F")},{a},">="&{start},{a},"<"&{end})'
PR = lambda col: f"Produits!${col}${P_FIRST}:${col}${P_LAST}"
kpis = [
    ("Le stock aujourd'hui", None, None),
    ("Nombre de produits", f"=COUNTA({PR('A')})", "int"),
    ('="Valeur du stock ("&$C$3&", prix d\'achat)"', f"=SUM({PR('L')})", "money"),
    ("Produits à commander", f'=COUNTIF({PR("M")},"À commander")', "int"),
    ("Produits en rupture", f'=COUNTIF({PR("M")},"Rupture")', "int"),
    ("Le mois affiché", None, None),
    ('="Ventes ("&$C$3&")"', "=" + month_sum("G", "Sortie") + cmd_month("H"), "money"),
    ('="Achats de marchandise ("&$C$3&")"', "=" + month_sum("G", "Entrée") + arr_month(), "money"),
    ('="Coût d\'achat des produits vendus ("&$C$3&")"', "=" + month_sum("H", None) + cmd_month("I"), "money"),
    ('="Marge brute ("&$C$3&")"', "=C12-C14", "money"),
    ("Taux de marge", '=IF(C12=0,"-",C15/C12)', "pct"),
]
r = 6
for label, formula, kind in kpis:
    lc = tb.cell(row=r, column=2, value=label)
    if formula is None:
        lc.font = font(size=12, bold=True, color=NAVY); r += 1; continue
    vc = tb.cell(row=r, column=3, value=formula)
    lc.font = font(size=11); lc.border = BORDER
    vc.font = font(size=12, bold=True); vc.fill = AUTO_FILL; vc.border = BORDER
    vc.alignment = Alignment(horizontal="right")
    vc.number_format = {"int": "0", "money": NUM, "pct": "0.0%"}[kind]
    tb.row_dimensions[r].height = 20
    r += 1
tb["D9"] = "La liste est dans l'onglet Produits (colonne État en orange)"
tb["D10"] = "La liste est dans l'onglet Produits (colonne État en rouge)"
tb["D15"] = "Ventes moins le coût d'achat de ce qui a été vendu (hors loyer, transport, salaires)"
for a in ("D9", "D10", "D15"): tb[a].font = font(size=9, italic=True, color="7F7F7F"); tb[a].alignment = Alignment(wrap_text=True)
tb.conditional_formatting.add("C9", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="FCD9A8")))
tb.conditional_formatting.add("C10", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F4B6B6")))

if AUTO:
    tb["B17"] = "Commandes reçues par le formulaire (mois)"; tb["B17"].font = font(size=11); tb["B17"].border = BORDER
    a = CM % ("A", "A")
    tb["C17"] = f'=COUNTIFS({a},">="&{start},{a},"<"&{end},{CM % ("J","J")},"<>Annulée")'
    tb["C17"].font = font(size=12, bold=True); tb["C17"].fill = AUTO_FILL; tb["C17"].border = BORDER; tb["C17"].number_format = "0"
    cm = wb.create_sheet("Commandes", 3)
    ccols = [("Date", 17), ("Client", 20), ("Téléphone", 15), ("Code produit", 12), ("Produit", 26), ("Quantité", 10),
             ("Prix unitaire", 12), ("Montant", 12), ("Coût d'achat", 12), ("Statut", 12), ("Adresse", 24), ("Note", 26)]
    headers(cm, 1, [(n, w, "x") for n, w in ccols])
    for r in range(2, 302):
        for ci in range(1, len(ccols) + 1):
            c = cm.cell(row=r, column=ci); c.font = font(); c.border = BORDER
            if ci == 1: c.number_format = "DD/MM/YYYY HH:MM"
            if ci in (6, 7, 8, 9): c.number_format = NUM
    dv_st = DataValidation(type="list", formula1='"Nouvelle,Livrée,Payée,Annulée"', allow_blank=True)
    cm.add_data_validation(dv_st); dv_st.add(f"J2:J{C_LAST}")
    cm.freeze_panes = "A2"
    cm.sheet_properties.tabColor = "7030A0"
    ar = wb.create_sheet("Arrivages", 4)
    acols = [("Date", 17), ("Code produit", 12), ("Produit", 26), ("Quantité", 10), ("Prix d'achat unitaire", 14),
             ("Montant", 13), ("Fournisseur", 22), ("Note", 26)]
    headers(ar, 1, [(n, w, "x") for n, w in acols])
    for r in range(2, 302):
        for ci in range(1, len(acols) + 1):
            c = ar.cell(row=r, column=ci); c.font = font(); c.border = BORDER
            if ci == 1: c.number_format = "DD/MM/YYYY HH:MM"
            if ci in (4, 5, 6): c.number_format = NUM
    ar.freeze_panes = "A2"
    ar.sheet_properties.tabColor = "1E7A46"
    mer = wb["Mode d'emploi"]
    last = mer.max_row + 2
    extra = [("h", "Onglet Commandes (automatique)"),
             ("t", "Les commandes saisies dans le formulaire en ligne arrivent ici toutes seules. N'ajoutez pas de ligne à la main, ne changez pas les titres de colonnes."),
             ("t", "Vous pouvez seulement changer la colonne Statut : Nouvelle, Livrée, Payée ou Annulée. Une commande Annulée ne compte plus dans le stock ni dans les ventes."),
             ("h", "Onglet Arrivages (automatique)"),
             ("t", "Chaque marchandise reçue et saisie dans le formulaire « Arrivage » arrive ici toute seule et augmente le stock. N'ajoutez pas de ligne à la main.")]
    for kind, txt in extra:
        c = mer.cell(row=last, column=2, value=txt); c.alignment = Alignment(wrap_text=True)
        c.font = font(size=12, bold=True, color=NAVY) if kind == "h" else font()
        last += 1
wb.save(OUT)
print("ok")
