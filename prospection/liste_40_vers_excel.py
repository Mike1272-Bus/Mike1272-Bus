"""Transforme liste_40_nouvelles.md en fichier Excel de suivi de prospection."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

src = open("liste_40_nouvelles.md", encoding="utf-8").read()
rows, cat = [], None
for line in src.splitlines():
    if line.startswith("## ") and not line.startswith("## Nos offres") and not line.startswith("## Le premier"):
        cat = line[3:].split(" :")[0].strip()
    m = re.match(r"^\| (\d+) \| (.+?) \| (.+?) \| (.+?) \| (.+?) \|$", line)
    if m and cat:
        num, ent, contact, offre, source = m.groups()
        link = re.search(r"\]\((https?://[^)]+)\)", source)
        label = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", source)
        rows.append([int(num), cat, ent, contact, offre, label, link.group(1) if link else ""])
assert len(rows) == 40, len(rows)

NAVY, YEL, GREY = "1B1F5E", "F5C518", "F2F4F8"
thin = Side(style="thin", color="D6DBE6")
wb = Workbook()
ws = wb.active
ws.title = "Prospects"
head = ["#", "Catégorie", "Entreprise", "Contact public", "Offre conseillée", "Source", "Lien source",
        "Numéro vérifié", "Date 1er message", "Réponse reçue", "Relance 1", "Relance 2", "Statut", "Notes"]
ws.append(head)
for r in rows:
    ws.append(r + ["Non", None, "", None, None, "À contacter", ""])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 32
widths = [5, 22, 38, 26, 34, 22, 34, 14, 15, 30, 13, 13, 16, 30]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w
n = len(rows) + 1
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    if row[0].row % 2 == 0:
        for c in row: c.fill = PatternFill("solid", fgColor=GREY)
    lk = row[6]
    if lk.value:
        lk.hyperlink = lk.value; lk.font = Font(color="1A73E8", underline="single")
    for c in (row[8], row[10], row[11]):
        c.number_format = "DD/MM/YYYY"
dv1 = DataValidation(type="list", formula1='"Oui,Non,Faux numéro"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"À contacter,Message envoyé,Relancé,Intéressé,Démo envoyée,Client,Pas intéressé"', allow_blank=True)
ws.add_data_validation(dv1); ws.add_data_validation(dv2)
dv1.add(f"H2:H{n}"); dv2.add(f"M2:M{n}")
ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=[f'$M2="Client"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=[f'$M2="Intéressé"'], fill=PatternFill("solid", fgColor="FFF2B3")))
ws.conditional_formatting.add(f"M2:M{n}", FormulaRule(formula=[f'$M2="Pas intéressé"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:N{n}"

# Résumé
s = wb.create_sheet("Résumé")
s["A1"] = "Suivi de la prospection : 40 nouvelles entreprises"; s["A1"].font = Font(bold=True, size=14, color=NAVY)
statuts = ["À contacter", "Message envoyé", "Relancé", "Intéressé", "Démo envoyée", "Client", "Pas intéressé"]
s["A3"], s["B3"] = "Statut", "Nombre"
for c in (s["A3"], s["B3"]):
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
for i, st in enumerate(statuts, 4):
    s.cell(i, 1, st); s.cell(i, 2, f'=COUNTIF(Prospects!M:M,"{st}")')
k = 4 + len(statuts)
s.cell(k + 1, 1, "Numéros vérifiés").font = Font(bold=True); s.cell(k + 1, 2, '=COUNTIF(Prospects!H:H,"Oui")')
s.cell(k + 3, 1, "Catégorie").font = Font(bold=True); s.cell(k + 3, 2, "Entreprises").font = Font(bold=True)
cats = list(dict.fromkeys(r[1] for r in rows))
for j, ct in enumerate(cats, k + 4):
    s.cell(j, 1, ct); s.cell(j, 2, f'=COUNTIF(Prospects!B:B,"{ct}")')
s.column_dimensions["A"].width = 40; s.column_dimensions["B"].width = 14

# Message et offres
msg = wb.create_sheet("Message et offres")
offres = re.findall(r"^\d\. \*\*(.+?)\*\*(.*)$", src.split("## Nos offres")[1].split("##")[0], re.M)
msg["A1"] = "Premier message (avec la question de diagnostic)"; msg["A1"].font = Font(bold=True, size=13, color=NAVY)
quote = src.split("## Le premier message")[1].split("\n", 1)[1]
lines = [l[2:] if l.startswith("> ") else l.lstrip(">") for l in quote.splitlines() if l.startswith(">")]
text = "\n".join(l for l in lines).strip()
msg["A2"] = text; msg["A2"].alignment = Alignment(wrap_text=True, vertical="top"); msg.row_dimensions[2].height = 150
msg["A4"] = "Nos offres"; msg["A4"].font = Font(bold=True, size=13, color=NAVY)
for i, (t, rest) in enumerate(offres, 5):
    msg.cell(i, 1, f"{i - 4}. {t}{rest}")
msg.cell(5 + len(offres) + 1, 1, "Avant d'écrire : vérifier que chaque numéro a bien un compte WhatsApp au nom de l'entreprise.").font = Font(italic=True)
msg.column_dimensions["A"].width = 110
wb.save("prospects_40_kinshasa.xlsx")
print("ok", len(rows))
